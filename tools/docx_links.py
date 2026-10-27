#!/usr/bin/env python3
"""Turn markdown links [text](url) that build_sheet.py leaves as plain text into
real, clickable hyperlinks (blue, underlined) in a .docx, including table cells.

Usage: python3 tools/docx_links.py sheet.docx
"""
import copy, re, sys
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+)\)")


def paragraphs(doc):
    yield from doc.paragraphs
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def make_run(rpr, text, link=False):
    r = OxmlElement("w:r")
    new = copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")
    if link:
        c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); new.append(c)
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); new.append(u)
    r.append(new)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    return r


def main(path):
    doc = Document(path)
    n = 0
    for p in paragraphs(doc):
        for run in list(p.runs):
            if not LINK.search(run.text):
                continue
            rpr = run._r.find(qn("w:rPr"))
            parts, pos = [], 0
            for m in LINK.finditer(run.text):
                if m.start() > pos:
                    parts.append(make_run(rpr, run.text[pos:m.start()]))
                h = OxmlElement("w:hyperlink")
                h.set(qn("r:id"), p.part.relate_to(m.group(2), RT.HYPERLINK, is_external=True))
                h.append(make_run(rpr, m.group(1), link=True))
                parts.append(h)
                pos = m.end(); n += 1
            if pos < len(run.text):
                parts.append(make_run(rpr, run.text[pos:]))
            for el in parts:
                run._r.addprevious(el)
            run._r.getparent().remove(run._r)
    doc.save(path)
    print(f"{n} links")


if __name__ == "__main__":
    main(sys.argv[1])
