#!/usr/bin/env python3
"""bible_docx_to_html.py — turn a bible .docx built by rossen-script-writer's
build_bible.py into HTML for the Drive upload (Drive converts it to a Google Doc).

Mirrors the .docx paragraph for paragraph, so the Google Doc keeps house format:
blank paragraph between every line, 1.15 line spacing, Arial, 23pt headers,
18pt body, bold red FF0000 cues, blue underlined links.

Usage: python3 tools/bible_docx_to_html.py "10_07 LIVE BIBLE.docx" out.html
"""
import html, sys
from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run

doc = Document(sys.argv[1])
rels = doc.part.rels
out = ['<html><head><meta charset="utf-8"></head><body style="font-family:Arial;font-size:18pt">']
P = "margin:0;line-height:1.15"
for par in doc.paragraphs:
    runs = [Run(c, par) for c in par._p if c.tag == qn("w:r")]
    sizes = {r.font.size.pt for r in runs if r.text.strip() and r.font.size}
    psize = max(sizes) if sizes else 18
    parts = []
    for child in par._p:
        if child.tag == qn("w:r"):
            r = Run(child, par)
            if not r.text:
                continue
            t = html.escape(r.text, quote=False)
            st = []
            if r.font.size and r.font.size.pt != psize:
                st.append(f"font-size:{r.font.size.pt:g}pt")
            if r.font.color is not None and r.font.color.type is not None and str(r.font.color.rgb) != "000000":
                st.append(f"color:#{r.font.color.rgb}")
            if st:
                t = f'<span style="{";".join(st)}">{t}</span>'
            if r.font.bold:
                t = f"<b>{t}</b>"
            parts.append(t)
        elif child.tag == qn("w:hyperlink"):
            url = rels[child.get(qn("r:id"))].target_ref
            text = "".join(t.text or "" for t in child.iter(qn("w:t")))
            parts.append(f'<a href="{html.escape(url)}" style="color:#1155CC">{html.escape(text, quote=False)}</a>')
    body = "".join(parts)
    fs = "" if psize == 18 else f";font-size:{psize:g}pt"
    out.append(f'<p style="{P}{fs}">{body or "&nbsp;"}</p>')
out.append("</body></html>")
open(sys.argv[2], "w", encoding="utf-8").write("\n".join(out))
print(f"wrote {sys.argv[2]}: {len(doc.paragraphs)} paragraphs")
