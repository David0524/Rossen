#!/usr/bin/env python3
"""build_source_log.py — render the companion source log for a Rossen bible.

Usage:
    python3 scripts/build_source_log.py sources.md "A-STORY HEADLINE - WED 10_14 - SOURCE LOG.docx"

Input is markdown with two sections, each holding one pipe table:

    ## CLAIMS
    | Claim | Status | Source | Scope |
    ## CLIPS
    | Beat | Orientation | Role | Candidate URL | Timecode | Known from |

Any other `## ` section (e.g. `## NEEDS A HUMAN`) is rendered as plain lines.
Validation errors exit 1 before anything is written:
  - Status outside CONFIRMED / PARTIALLY CONFIRMED / UNVERIFIED / CONTRADICTED
  - empty Source or Scope on a claim; empty Candidate URL or Known from on a clip
  - missing CLAIMS section
Requires python-docx (pip install python-docx --break-system-packages).
"""
import re
import sys

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx missing: pip install python-docx --break-system-packages")

STATUSES = {"CONFIRMED", "PARTIALLY CONFIRMED", "UNVERIFIED", "CONTRADICTED"}
STATUS_COLOR = {"CONFIRMED": "1E7B34", "PARTIALLY CONFIRMED": "B26B00",
                "UNVERIFIED": "B26B00", "CONTRADICTED": "FF0000"}
REQUIRED = {"CLAIMS": ["Source", "Scope"], "CLIPS": ["Candidate URL", "Known from"]}


def parse(path):
    sections, cur = {}, None
    for raw in open(path, encoding="utf-8"):
        line = raw.rstrip("\n")
        m = re.match(r"^##\s+(.+?)\s*$", line)
        if m:
            cur = m.group(1).strip().upper()
            sections[cur] = {"rows": [], "text": []}
            continue
        if cur is None:
            continue
        if line.strip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                continue
            sections[cur]["rows"].append(cells)
        elif line.strip():
            sections[cur]["text"].append(line.strip())
    return sections


def validate(sections):
    errs = []
    if "CLAIMS" not in sections or len(sections["CLAIMS"]["rows"]) < 2:
        errs.append("no CLAIMS table with at least one row")
    for name, req in REQUIRED.items():
        if name not in sections or not sections[name]["rows"]:
            continue
        head, *rows = sections[name]["rows"]
        idx = {h: i for i, h in enumerate(head)}
        for col in req:
            if col not in idx:
                errs.append(f"{name}: missing column {col!r}")
        for n, row in enumerate(rows, 1):
            row = row + [""] * (len(head) - len(row))
            for col in req:
                if col in idx and not row[idx[col]]:
                    errs.append(f"{name} row {n}: empty {col!r}")
            if name == "CLAIMS" and "Status" in idx:
                st = row[idx["Status"]].upper()
                if st not in STATUSES:
                    errs.append(f"CLAIMS row {n}: status {row[idx['Status']]!r} "
                                f"not one of {sorted(STATUSES)}")
    return errs


def run(par, text, bold=False, color=None, size=10):
    r = par.add_run(text)
    r.font.name = "Arial"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def build(sections, out, title):
    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Inches(11), Inches(8.5)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Inches(0.6))
    run(doc.add_paragraph(), title, bold=True, size=16)
    order = ["CLAIMS", "CLIPS"] + [k for k in sections if k not in ("CLAIMS", "CLIPS")]
    for name in order:
        if name not in sections:
            continue
        s = sections[name]
        run(doc.add_paragraph(), name, bold=True, size=13)
        for t in s["text"]:
            run(doc.add_paragraph(), re.sub(r"^[-*]\s*", "", t))
        if not s["rows"]:
            continue
        head, *rows = s["rows"]
        tbl = doc.add_table(rows=1, cols=len(head))
        tbl.style = "Table Grid"
        for i, h in enumerate(head):
            c = tbl.rows[0].cells[i]
            c.text = ""
            run(c.paragraphs[0], h, bold=True)
            shade(c, "D9D9D9")
        st_i = head.index("Status") if "Status" in head else None
        for row in rows:
            row = row + [""] * (len(head) - len(row))
            cells = tbl.add_row().cells
            for i, v in enumerate(row[:len(head)]):
                cells[i].text = ""
                color = STATUS_COLOR.get(v.upper()) if i == st_i else None
                run(cells[i].paragraphs[0], v, bold=bool(color), color=color)
        doc.add_paragraph()
    doc.save(out)


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    src, out = sys.argv[1], sys.argv[2]
    sections = parse(src)
    errs = validate(sections)
    if errs:
        print("source log NOT written:")
        for e in errs:
            print(f"  x {e}")
        sys.exit(1)
    title = re.sub(r"\.docx$", "", out.rsplit("/", 1)[-1], flags=re.I)
    build(sections, out, title)
    n_claims = len(sections["CLAIMS"]["rows"]) - 1
    n_clips = max(0, len(sections.get("CLIPS", {"rows": []})["rows"]) - 1)
    print(f"wrote {out}  ({n_claims} claims, {n_clips} clips)")


if __name__ == "__main__":
    main()
