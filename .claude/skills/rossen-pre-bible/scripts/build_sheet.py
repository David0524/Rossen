#!/usr/bin/env python3
"""Render a Pre-Bible Pitch Sheet markdown source into a .docx.

Usage:
    python3 build_sheet.py sheet.md -o "Pre_Bible_Pitch_Sheet.docx"

Source conventions this expects:
    # TITLE                      -> document title (first H1 only)
    ## SECTION                   -> major section (shortlist, back matter)
    # N. STORY HEADLINE          -> a story page; forces a page break before it
    ### Subsection               -> story page subsection
    | a | b |                    -> pipe table (the shortlist)
    > quoted line                -> blockquote (teases, viral callouts)
    - bullet                     -> bullet
    **bold** and *italic*        -> inline emphasis

Story pages are detected as H1s beginning with a digit. Each gets a hard page
break before it. The shortlist stays on page one.

python-docx cannot report where a page actually breaks. After building, convert
to PDF and look at the pages.
"""

import argparse
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor

STORY_H1 = re.compile(r"^#\s+(\d+)\.\s+(.*)$")
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")


def add_runs(paragraph, text):
    """Split inline **bold** / *italic* markers into runs."""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            paragraph.add_run(chunk[1:-1]).italic = True
        else:
            paragraph.add_run(chunk)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    for attr in ("top_margin", "bottom_margin"):
        setattr(sec, attr, Inches(0.7))
    for attr in ("left_margin", "right_margin"):
        setattr(sec, attr, Inches(0.8))
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(4)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Column widths tuned for: # | Story | Sexy beat | Viral | Slot | Confidence
    defaults = [0.35, 1.45, 2.35, 0.85, 1.15, 0.85]
    widths = defaults if len(header) == 6 else [6.9 / len(header)] * len(header)
    for cell, text, w in zip(table.rows[0].cells, header, widths):
        cell.width = Inches(w)
        p = cell.paragraphs[0]
        run = p.add_run(re.sub(r"\*", "", text))
        run.bold = True
        run.font.size = Pt(9.5)
    for row in body:
        cells = table.add_row().cells
        for cell, text, w in zip(cells, row, widths):
            cell.width = Inches(w)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, text)
            for run in p.runs:
                run.font.size = Pt(9.5)
    doc.add_paragraph()


def render(md, out):
    doc = Document()
    setup(doc)

    lines = md.splitlines()
    i = 0
    seen_title = False
    pending_table = []

    def flush_table():
        nonlocal pending_table
        if pending_table:
            build_table(doc, pending_table)
            pending_table = []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        # --- tables -------------------------------------------------------
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                pending_table.append(cells)
            i += 1
            continue
        flush_table()

        if not stripped:
            i += 1
            continue

        # --- headings -----------------------------------------------------
        m = STORY_H1.match(stripped)
        if m:
            page_break(doc)
            h = doc.add_heading(f"{m.group(1)}. {m.group(2)}", level=1)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.size = Pt(16)
            i += 1
            continue

        if stripped.startswith("### "):
            h = doc.add_heading(stripped[4:], level=3)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
                run.font.size = Pt(11)
            i += 1
            continue

        if stripped.startswith("## "):
            h = doc.add_heading(stripped[3:], level=2)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.font.size = Pt(13)
            i += 1
            continue

        if stripped.startswith("# "):
            text = stripped[2:]
            if not seen_title:
                h = doc.add_heading(text, level=0)
                seen_title = True
            else:
                page_break(doc)
                h = doc.add_heading(text, level=1)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
            i += 1
            continue

        # --- blockquote ---------------------------------------------------
        if stripped.startswith(">"):
            block = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                block.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, " ".join(x for x in block if x))
            for run in p.runs:
                run.italic = True
            continue

        # --- bullets ------------------------------------------------------
        if stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, stripped[2:])
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, re.sub(r"^\d+\.\s", "", stripped))
            i += 1
            continue

        if stripped.startswith("<!--"):
            i += 1
            continue

        # --- prose --------------------------------------------------------
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    flush_table()
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source")
    ap.add_argument("-o", "--output", default="Pre_Bible_Pitch_Sheet.docx")
    args = ap.parse_args()

    with open(args.source, encoding="utf-8") as fh:
        md = fh.read()

    path = render(md, args.output)
    print(f"wrote {path}")
    print("Now convert to PDF and LOOK at the pages — one page per story is only")
    print("real if you have seen it:")
    print(f'  soffice --headless --convert-to pdf "{path}"')


if __name__ == "__main__":
    sys.exit(main())
